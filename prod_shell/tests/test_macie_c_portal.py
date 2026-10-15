"""MACIE C authenticated browser acceptance tests.

Tests run real Flask routing and real session/CSRF logic, but stub AI
response to avoid network/secret dependency. Independent live-provider
validation remains a separate acceptance gate.
"""
import hashlib
import io
import json
import os
import time

import pytest

import macie_web
from macie_web import app
from prod_shell.src import macie_c_portal as portal
from prod_shell.src import customer_openrouter_keys as customer_keys


def password_record(password: str, salt: bytes = b"abcdefghijklmnop"):
    rounds = 260000
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, rounds)
    return f"pbkdf2_sha256${rounds}${salt.hex()}${dk.hex()}"


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setenv("MACIE_C_SESSION_SECRET", "session-secret-for-tests-only-" * 3)
    monkeypatch.setenv("MACIE_C_ENABLED", "1")
    monkeypatch.setenv("MACIE_C_USERS_JSON", json.dumps({
        "owner@example.com": password_record("owner-pilot-strong-password"),
        "tester@example.com": password_record("tester-pilot-strong-password"),
    }))
    monkeypatch.delenv("MACIE_ADMIN_KEY", raising=False)
    monkeypatch.setattr(customer_keys, "verify_openrouter_key", lambda key: None)
    customer_keys.clear_for_testing()
    app.config.update(TESTING=True, SECRET_KEY="test-secret-not-live-" * 4)
    with portal._lock:
        portal._jobs.clear()
        portal._failed_logins.clear()
    with app.test_client() as client:
        yield client


def sign_in(client, email="owner@example.com", password="owner-pilot-strong-password"):
    r = client.post("/macie/c/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.json
    with client.session_transaction() as s:
        return s["macie_c_csrf"]


def sample_key():
    return "sk-or-" + "demo-" + ("testing" * 7)


def _session_id(client):
    with client.session_transaction() as sess:
        return sess["macie_c_sid"]


def test_disabled_by_default_if_configuration_missing(client, monkeypatch):
    monkeypatch.setenv("MACIE_C_ENABLED", "0")
    r = client.get("/macie/c")
    assert r.status_code == 503
    assert r.headers["Cache-Control"].startswith("no-store")


def test_signin_and_protected_portal(client):
    home = client.get("/macie/c")
    assert home.status_code == 200
    assert b"Sign in" in home.data
    assert b"Run independent MACIE C review" not in home.data
    assert home.headers.get("Content-Security-Policy")
    assert "HttpOnly" not in home.headers.get("Set-Cookie", "")
    unauthorized = client.post("/macie/c/review", json={"document": "hi"})
    assert unauthorized.status_code == 401
    assert client.post("/macie/c/login", json={
        "email": "owner@example.com", "password": "incorrect"
    }).status_code == 401
    csrf = sign_in(client)
    protected = client.get("/macie/c")
    assert protected.status_code == 200
    assert b"Run independent MACIE C review" in protected.data
    assert "HttpOnly" in client.post("/macie/c/login", json={
        "email": "owner@example.com", "password": "owner-pilot-strong-password"
    }).headers.get("Set-Cookie", "")
    assert isinstance(csrf, str)


def test_legacy_paid_query_requires_admin_auth(client):
    r = client.post("/macie/query", json={"query": "expensive API question"})
    assert r.status_code == 401
    r = client.get("/macie")
    assert r.status_code == 302
    assert "/macie/a" in r.headers["Location"]
    root = client.get("/")
    assert root.status_code == 302
    assert "/macie/a" in root.headers["Location"]


def test_verified_cloudflare_owner_restores_private_macie_a(client, monkeypatch):
    monkeypatch.setattr(
        macie_web,
        "_cloudflare_access_email",
        lambda: "peteblozis@gmail.com",
    )

    home = client.get("/macie")
    assert home.status_code == 200
    assert b"Multi-AI" in home.data
    assert b"Ask MACIE" in home.data

    alias = client.get("/macie/a")
    assert alias.status_code == 200

    # The verified Access identity passes the paid-query authentication
    # boundary. Empty input stops before any provider call is made.
    query = client.post("/macie/query", json={})
    assert query.status_code == 400
    assert query.json["error"] == "No query provided"


def test_csrf_and_origin_rejected(client):
    csrf = sign_in(client)
    no_csrf = client.post("/macie/c/review", json={"document": "Something"})
    assert no_csrf.status_code == 403
    wrong_origin = client.post("/macie/c/review",
        json={"document": "Something"},
        headers={"X-CSRF-Token": csrf, "Origin": "https://evil.example"})
    assert wrong_origin.status_code == 403


def test_real_web_submission_queue_and_owner_separation(client, monkeypatch):
    calls = []
    def fake_review(doc, goal, options=None, *, customer_openrouter_key=None):
        calls.append((doc, goal))
        return {
            "edition": "MACIE C",
            "status": "CONSENSUS_WITH_DIVERGENCES",
            "synthesized_review": "Check all marketing statistics and never guarantee sales outcomes.",
            "divergences": [{"claim": "brochure CTA", "positions": {
                "claude": "call your agent", "chatgpt": "request a consultation",
            }}],
            "model_outputs": [],
            "models_succeeded": ["chatgpt", "claude"],
            "confidence": "medium",
        }
    monkeypatch.setattr(portal, "review_document", fake_review)
    csrf = sign_in(client)
    customer_keys.save_key("owner@example.com", _session_id(client), sample_key())
    submitted = client.post("/macie/c/review", data={
        "document": "Copilot claims that 73% faster sales are guaranteed.",
        "goal": "Review for HOMEstretch and make a stronger brochure.",
    }, headers={"X-CSRF-Token": csrf})
    assert submitted.status_code == 202, submitted.json
    jid = submitted.json["job_id"]
    assert jid and submitted.json["status"] == "queued"
    for _ in range(30):
        report = client.get("/macie/c/review/" + jid)
        if report.json["status"] == "finished":
            break
        time.sleep(.03)
    assert report.status_code == 200
    assert report.json["status"] == "finished"
    assert report.json["result"]["status"] == "CONSENSUS_WITH_DIVERGENCES"
    assert "73%" in calls[0][0]

    other = app.test_client()
    sign_in(other, "tester@example.com", "tester-pilot-strong-password")
    forbidden = other.get("/macie/c/review/" + jid)
    assert forbidden.status_code == 404


def test_oversized_submission_rejected(client):
    csrf = sign_in(client)
    r = client.post("/macie/c/review", json={"document": "x" * 50001},
                    headers={"X-CSRF-Token": csrf})
    assert r.status_code == 400


def test_invalid_file_extension_rejected(client):
    csrf = sign_in(client)
    r = client.post("/macie/c/review",
        data={"file": (io.BytesIO(b"not safe"), "danger.exe"), "goal": "Review"},
        content_type="multipart/form-data", headers={"X-CSRF-Token": csrf})
    assert r.status_code == 400
    assert "docx" in r.json["error"]


def test_valid_docx_extraction(client, monkeypatch):
    from docx import Document
    d = Document()
    d.add_paragraph("Important Realtor brochure proposal.")
    d.add_table(rows=1, cols=2).cell(0, 0).text = "HOMEstretch"
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    def fake_review(doc, goal, options=None, *, customer_openrouter_key=None):
        assert "Important Realtor brochure proposal." in doc
        assert "HOMEstretch" in doc
        return {"status": "NON-CONSENSUS", "models_succeeded": [],
                "confidence": "low", "model_outputs": [],
                "synthesized_review": "", "raw_fallback": ""}
    monkeypatch.setattr(portal, "review_document", fake_review)
    csrf = sign_in(client)
    customer_keys.save_key("owner@example.com", _session_id(client), sample_key())
    r = client.post("/macie/c/review", data={
        "file": (buf, "sample.docx"), "goal": "Review",
    }, content_type="multipart/form-data", headers={"X-CSRF-Token": csrf})
    assert r.status_code == 202


def test_login_bruteforce_limit(client):
    for i in range(6):
        assert client.post("/macie/c/login", json={
            "email": "owner@example.com", "password": "bad"
        }).status_code == 401
    assert client.post("/macie/c/login", json={
        "email": "owner@example.com", "password": "bad"
    }).status_code == 429


def test_logout_revokes_session(client):
    csrf = sign_in(client)
    result = client.post("/macie/c/logout", headers={"X-CSRF-Token":csrf})
    assert result.status_code == 200
    assert client.post("/macie/c/review",json={"document":"x"}).status_code == 401


def test_buyers_connect_and_revoke_own_key_through_browser_api(client):
    csrf=sign_in(client)
    missing=client.post("/macie/c/review",json={"document":"Public draft"},
                        headers={"X-CSRF-Token":csrf})
    assert missing.status_code==409
    response=client.post("/macie/c/provider",json={"api_key":sample_key()},
                         headers={"X-CSRF-Token":csrf})
    assert response.status_code==200
    assert response.json["connected"]
    assert sample_key() not in response.get_data(as_text=True)
    assert sample_key() not in str(response.headers)
    status=client.get("/macie/c/provider")
    assert status.status_code==200 and status.json["connected"]
    assert sample_key() not in json.dumps(status.json)
    assert sample_key() not in str(client.get_cookie("session"))
    other=app.test_client()
    sign_in(other,"tester@example.com","tester-pilot-strong-password")
    assert not other.get("/macie/c/provider").json["connected"]
    removal=client.delete("/macie/c/provider",headers={"X-CSRF-Token":csrf})
    assert removal.status_code==200
    assert not client.get("/macie/c/provider").json["connected"]


def test_provider_connection_csrf_and_origin(client):
    assert client.post("/macie/c/provider",json={"api_key":sample_key()}).status_code==401
    csrf=sign_in(client)
    assert client.post("/macie/c/provider",json={"api_key":sample_key()}).status_code==403
    assert client.post("/macie/c/provider",json={"api_key":sample_key()},
        headers={"X-CSRF-Token":csrf,"Origin":"https://evil.example"}).status_code==403
    bad=client.post("/macie/c/provider",json={"api_key":"invalid"},
        headers={"X-CSRF-Token":csrf})
    assert bad.status_code==400


def test_logout_erases_buyer_key_and_blocks_further_provider_use(client):
    csrf=sign_in(client)
    r=client.post("/macie/c/provider",json={"api_key":sample_key()},
                  headers={"X-CSRF-Token":csrf})
    assert r.status_code==200
    sid=_session_id(client)
    logout=client.post("/macie/c/logout",headers={"X-CSRF-Token":csrf})
    assert logout.status_code==200
    assert customer_keys.get_key("owner@example.com",sid) is None
    assert client.get("/macie/c/provider").status_code==401

"""End-to-end autonomous commercial MACIE C API tests.

All users, emails, processor events, and model responses are isolated test
fixtures. No payment, key, third-party model or SMTP account is accessed.
The commercial routes stay disabled in production until release gates pass.
"""
import json
import io
import time
import tempfile
from pathlib import Path
from sqlalchemy import create_engine

import pytest

from macie_web import app
from prod_shell.src import commercial_portal as portal
from prod_shell.src import commercial_store as db


def buyer_key(letter="Q"):
    return "sk-or-v1-" + letter * 42


@pytest.fixture
def fixture(monkeypatch):
    # Use a REAL temporary SQLite database file so in-process worker
    # threads get independent database connections. In-memory StaticPool
    # shares a single connection between threads and is not concurrency-safe.
    with tempfile.TemporaryDirectory(prefix="macie-c-commercial-ci-") as tmp:
        engine=create_engine(
            "sqlite+pysqlite:///"+str(Path(tmp)/"commercial_test.sqlite"),
            connect_args={"timeout":15},future=True,
        )
        store=db.Store(engine,"CI-only-commercial-identity-vault-test-0000000000000")
        store.initialize_for_tests_only()
        yield from _test_fixture(monkeypatch,store)
        engine.dispose()


def _test_fixture(monkeypatch,store):
    messages=[]
    monkeypatch.setattr(portal,"_configured",lambda:True)
    monkeypatch.setattr(portal,"_store",lambda:store)
    monkeypatch.setattr(portal,"send_token_message",lambda email,token,purpose:messages.append((email,token,purpose)))
    monkeypatch.setattr(portal,"verify_openrouter_key",lambda key: None)
    calls=[]
    def synthetic_model(document,goal,options=None,*,customer_openrouter_key=None):
        calls.append((document,goal,customer_openrouter_key))
        assert customer_openrouter_key and customer_openrouter_key.startswith("sk-or-")
        return {
            "edition":"MACIE C","status":"CONSENSUS_WITH_DIVERGENCES",
            "confidence":"high",
            "models_succeeded":["claude","chatgpt"],
            "model_outputs":[
                {"model_id":"claude","model_version":"fixture-only-claude","raw_review":"Verify statistics."},
                {"model_id":"chatgpt","model_version":"fixture-only-gpt","raw_review":"Respect user privacy."},
            ],
            "divergences":[{"claim":"Example difference","positions":{"claude":"A","chatgpt":"B"}}],
            "synthesized_review":"Test only: verify facts and privacy.",
            "live_acceptance_passed":False,
        }
    monkeypatch.setattr(portal,"review_document",synthetic_model)
    app.config.update(TESTING=True,SECRET_KEY="commercial-test-cookie-key-never-production-" * 2)
    with app.test_client() as client:
        yield {"client":client,"store":store,"emails":messages,"calls":calls}


def customer(fixture,email="buyer@local.test",company="Buyer Test Corporation"):
    cli=fixture["client"]
    sign=cli.post("/macie/c/commercial/signup",json={
        "email":email,"company":company,
        "password":"Safe-long-commercial-password-123",
        "accept_terms":True,
    })
    assert sign.status_code==202,sign.json
    assert "token" not in sign.get_data(as_text=True)
    token=fixture["emails"][-1][1]
    assert fixture["emails"][-1][2]=="verify"
    verify=cli.post("/macie/c/commercial/verify-email",json={"token":token})
    assert verify.status_code==200,verify.json
    assert verify.json["verified"]
    login=cli.post("/macie/c/commercial/login",json={
        "email":email,"password":"Safe-long-commercial-password-123",
    })
    assert login.status_code==200,login.json
    status=cli.get("/macie/c/commercial/status").json
    assert status["authenticated"]
    return {"email":email,"csrf":login.json["csrf"],
            "tenant":status["organization_id"],
            "user":fixture["store"].get_account_by_email(email)["id"]}


def paid(store,tenant,event="evt-purchase-1",limit=2):
    assert store.apply_verified_billing_event(
        event_id=event,event_type="invoice.paid",event_at=int(time.time()),
        tenant_id=tenant,subscription_status="active",
        customer_id="cus_"+tenant,subscription_id="sub_"+tenant,
        review_limit=limit,billing_verified=True,
    )


def test_signup_link_proves_email_and_enables_independent_signin(fixture):
    cli=fixture["client"]
    home=cli.get("/macie/c/commercial")
    assert home.status_code==200
    assert b"Create your account" in home.data
    member=customer(fixture)
    assert fixture["store"].get_account_by_email(member["email"])["verified"]
    status=cli.get("/macie/c/commercial/status").json
    assert status["subscription"]["status"]=="inactive"
    assert not status["subscription"]["billing_verified"]
    assert not status["provider_connected"]
    assert cli.post("/macie/c/commercial/checkout",json={},
                    headers={"X-CSRF-Token":member["csrf"]}).status_code==503


def test_commercial_account_cannot_spoof_billing_or_use_owner_key(fixture):
    cli=fixture["client"];member=customer(fixture)
    auth={"X-CSRF-Token":member["csrf"]}
    before=cli.post("/macie/c/commercial/review",json={
        "document":"A sample brochure","goal":"Improve it","paid":True,
        "tenant_id":"some-other-tenant",
    },headers=auth)
    assert before.status_code==403
    assert fixture["calls"]==[]
    assert "subscription" in before.json["error"]


def test_buyer_connects_their_key_and_runs_paid_review_independently(fixture):
    cli=fixture["client"];member=customer(fixture);h={"X-CSRF-Token":member["csrf"]}
    key=buyer_key()
    provider=cli.post("/macie/c/commercial/provider",json={"api_key":key},headers=h)
    assert provider.status_code==200,provider.json
    assert key not in provider.get_data(as_text=True)
    assert key not in str(cli.get_cookie("session"))
    paid(fixture["store"],member["tenant"])
    sent=cli.post("/macie/c/commercial/review",json={
        "document":"HOMEstretch claims 73% faster sales, without sourcing.",
        "goal":"Examine the marketing claims, then revise responsibly.",
    },headers=h)
    assert sent.status_code==202,sent.json
    job_id=sent.json["job_id"]
    for _ in range(100):
        status=cli.get("/macie/c/commercial/review/"+job_id)
        if status.json["status"]=="finished":
            assert status.json["result"]["status"]=="CONSENSUS_WITH_DIVERGENCES"
            assert "Test only" in status.json["result"]["synthesized_review"]
            break
        time.sleep(.03)
    else:
        pytest.fail("In-process commercial job did not complete.")
    assert len(fixture["calls"])==1
    assert fixture["calls"][0][2]==key
    assert fixture["store"].entitlement(member["tenant"]).reviews_consumed==1
    assert key not in status.get_data(as_text=True)


def test_customer_key_revocation_and_no_fallback(fixture):
    cli=fixture["client"];m=customer(fixture);h={"X-CSRF-Token":m["csrf"]}
    cli.post("/macie/c/commercial/provider",json={"api_key":buyer_key()},headers=h)
    paid(fixture["store"],m["tenant"])
    assert cli.delete("/macie/c/commercial/provider",headers=h).status_code==200
    r=cli.post("/macie/c/commercial/review",
               json={"document":"A short sample","goal":"Review"},headers=h)
    assert r.status_code==403
    assert "key" in r.json["error"]
    assert fixture["calls"]==[]


def test_unauthenticated_cross_origin_and_missing_csrf_refused(fixture):
    cli=fixture["client"]
    assert cli.get("/macie/c/commercial/status").json["authenticated"] is False
    assert cli.post("/macie/c/commercial/provider",json={"api_key":buyer_key()}).status_code==401
    m=customer(fixture)
    assert cli.post("/macie/c/commercial/provider",json={"api_key":buyer_key()}).status_code==403
    evil=cli.post("/macie/c/commercial/provider",json={"api_key":buyer_key()},
                  headers={"X-CSRF-Token":m["csrf"],"Origin":"https://evil.example"})
    assert evil.status_code==403
    assert cli.get("/macie/c/commercial").status_code==200


def test_review_result_isolation_between_two_customer_organizations(fixture):
    cli=fixture["client"];buyer1=customer(fixture)
    h={"X-CSRF-Token":buyer1["csrf"]}
    cli.post("/macie/c/commercial/provider",json={"api_key":buyer_key()},headers=h)
    paid(fixture["store"],buyer1["tenant"])
    review=cli.post("/macie/c/commercial/review",json={
        "document":"Private first tenant","goal":"Review for clarity"},headers=h)
    assert review.status_code==202
    job_id=review.json["job_id"]
    cli.post("/macie/c/commercial/logout",json={},headers=h)
    second=customer(fixture,email="other@local.test",company="Second Test Corporation")
    assert second["tenant"]!=buyer1["tenant"]
    assert cli.get("/macie/c/commercial/review/"+job_id).status_code==404
    assert not cli.get("/macie/c/commercial/provider").json["connected"]


def test_password_reset_is_customer_operated_and_revokes_old_password(fixture):
    cli=fixture["client"];m=customer(fixture)
    request=cli.post("/macie/c/commercial/request-reset",json={"email":m["email"]})
    assert request.status_code==202
    assert fixture["emails"][-1][2]=="reset"
    token=fixture["emails"][-1][1]
    update=cli.post("/macie/c/commercial/reset-password",json={
        "token":token,"new_password":"Another-longer-self-managed-passphrase",
    })
    assert update.status_code==200
    bad=cli.post("/macie/c/commercial/login",json={
        "email":m["email"],"password":"Safe-long-commercial-password-123",
    })
    assert bad.status_code==401
    fresh=cli.post("/macie/c/commercial/login",json={
        "email":m["email"],"password":"Another-longer-self-managed-passphrase",
    })
    assert fresh.status_code==200


def test_unverified_account_cannot_login_or_create_tenant(fixture):
    cli=fixture["client"]
    signup=cli.post("/macie/c/commercial/signup",json={
        "email":"pending@local.test","company":"No Verify Inc",
        "password":"Very-long-and-safe-password-for-signup",
        "accept_terms":True,
    })
    assert signup.status_code==202
    assert not fixture["store"].get_account_by_email("pending@local.test")["verified"]
    r=cli.post("/macie/c/commercial/login",json={
        "email":"pending@local.test","password":"Very-long-and-safe-password-for-signup",
    })
    assert r.status_code==401


def test_real_product_is_off_unless_all_operational_dependencies_exist(fixture,monkeypatch):
    monkeypatch.setattr(portal,"_configured",lambda:False)
    cli=fixture["client"]
    blocked=cli.get("/macie/c/commercial")
    assert blocked.status_code==503
    assert "not yet available" in blocked.get_data(as_text=True)
    assert cli.post("/macie/c/commercial/signup",json={
        "email":"any@local.test","password":"Strong-long-passphrase-123456",
        "company":"Company","accept_terms":True,
    }).status_code==503


def test_login_attempts_persist_and_throttle(fixture):
    cli=fixture["client"];m=customer(fixture)
    cli.post("/macie/c/commercial/logout",json={},
             headers={"X-CSRF-Token":m["csrf"]})
    for _ in range(6):
        result=cli.post("/macie/c/commercial/login",json={
            "email":m["email"],"password":"not-the-right-password",
        })
        assert result.status_code==401
    again=cli.post("/macie/c/commercial/login",json={
        "email":m["email"],"password":"Safe-long-commercial-password-123",
    })
    assert again.status_code==429


def test_paid_buyer_uploads_existing_word_document_without_owner_help(fixture):
    from docx import Document
    cli=fixture["client"]
    member=customer(fixture)
    headers={"X-CSRF-Token":member["csrf"]}
    assert cli.post("/macie/c/commercial/provider",
        json={"api_key":buyer_key()},headers=headers).status_code==200
    paid(fixture["store"],member["tenant"],event="evt-docx-buy",limit=2)
    doc=Document()
    doc.add_paragraph("Existing Realtor pitch for independent review.")
    doc.add_table(rows=1,cols=2).cell(0,0).text="Original brochure details"
    stream=io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    submitted=cli.post("/macie/c/commercial/review",data={
        "file":(stream,"brochure.docx"),
        "goal":"Identify unsupported claims and improve the existing copy",
    },content_type="multipart/form-data",headers=headers)
    assert submitted.status_code==202,submitted.json
    job_id=submitted.json["job_id"]
    for attempt in range(80):
        current=cli.get("/macie/c/commercial/review/"+job_id)
        if current.json["status"]=="finished":
            break
        time.sleep(.03)
    else:
        pytest.fail("Commercial Word document review did not complete.")
    assert len(fixture["calls"])==1
    assert "Existing Realtor pitch" in fixture["calls"][0][0]
    assert "Original brochure details" in fixture["calls"][0][0]
    assert fixture["store"].entitlement(member["tenant"]).reviews_consumed==1

"""MACIE C authenticated, nonpersistent document review portal.

The /macie/c routes are separate from private MACIE A. The portal accepts
pasted text and small DOCX/TXT files without editing Render configuration.
Results are held in process memory only for 30 minutes; use Download to
preserve a report. Customer-facing access is OFF until explicitly enabled.
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import secrets
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta

from flask import (
    Blueprint, abort, current_app, g, jsonify, make_response,
    render_template_string, request, session,
)

from prod_shell.src.macie_c import MAX_DOCUMENT_CHARS, review_document
from prod_shell.src.macie_c_template import PAGE_HTML


bp = Blueprint("macie_c_portal", __name__, url_prefix="/macie/c")

# One small Render worker: no external queue or durable storage; no documents
# or AI responses are written to local files or app logs by the portal.
_MAX_REQUEST_BYTES = 2_500_000
_MAX_FILE_BYTES = 2_000_000
_MAX_UNCOMPRESSED_ZIP_BYTES = 28_000_000
_JOB_TTL_SECONDS = 1800
_EXECUTOR = ThreadPoolExecutor(max_workers=2, thread_name_prefix="macie-c-review")
_lock = threading.RLock()
_jobs: dict[str, dict] = {}
_failed_logins: dict[str, list[float]] = {}


def configure_portal(app):
    """Install secure cookie settings and register the customer-only blueprint."""
    secret = os.environ.get("MACIE_C_SESSION_SECRET", "")
    if len(secret) >= 40:
        app.secret_key = secret
    app.config.update(
        SESSION_COOKIE_SECURE=True,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Strict",
        PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
        MAX_CONTENT_LENGTH=_MAX_REQUEST_BYTES,
    )
    app.register_blueprint(bp)


def _users() -> dict[str, str]:
    try:
        raw = json.loads(os.environ.get("MACIE_C_USERS_JSON", "{}"))
        if not isinstance(raw, dict):
            return {}
        return {str(k).strip().lower(): str(v) for k, v in raw.items()}
    except (ValueError, TypeError):
        return {}


def _is_ready() -> bool:
    return (
        os.environ.get("MACIE_C_ENABLED") == "1"
        and len(os.environ.get("MACIE_C_SESSION_SECRET", "")) >= 40
        and bool(_users())
    )


def _check_password(password: str, record: str) -> bool:
    try:
        scheme, rounds, salt_hex, digest_hex = record.split("$")
        if scheme != "pbkdf2_sha256":
            return False
        count = int(rounds)
        if count < 260000 or count > 1000000:
            return False
        salt = bytes.fromhex(salt_hex)
        digest = bytes.fromhex(digest_hex)
        if len(salt) < 16 or len(digest) != 32:
            return False
        computed = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, count)
        return hmac.compare_digest(computed, digest)
    except (ValueError, TypeError, UnicodeError):
        return False


def _current_user() -> str | None:
    email = session.get("macie_c_email")
    if isinstance(email, str) and email in _users():
        return email
    return None


def _same_origin() -> bool:
    origin = request.headers.get("Origin", "")
    if not origin:
        # CSRF token is still mandatory for authenticated writes.
        return True
    from urllib.parse import urlsplit
    parsed = urlsplit(origin)
    # Render terminates TLS in front of Gunicorn; Flask may see internal
    # HTTP even though the browser uses HTTPS. Compare browser-facing host.
    return (
        parsed.netloc.lower() == request.host.lower()
        and parsed.scheme == ("http" if current_app.testing else "https")
        and parsed.path in ("", "/")
    )


def _authenticated():
    if not _current_user():
        return jsonify({"error": "Sign in before submitting documents."}), 401
    return None


def _csrf_ok() -> bool:
    supplied = request.headers.get("X-CSRF-Token", "")
    stored = session.get("macie_c_csrf", "")
    return (
        bool(stored)
        and bool(supplied)
        and hmac.compare_digest(str(stored), str(supplied))
        and _same_origin()
    )


def _login_blocked(key: str) -> bool:
    now = time.monotonic()
    with _lock:
        attempts = [t for t in _failed_logins.get(key, []) if now - t < 900]
        _failed_logins[key] = attempts
        return len(attempts) >= 6


def _record_login_failure(key: str):
    with _lock:
        _failed_logins.setdefault(key, []).append(time.monotonic())


def _cleanup_jobs():
    now = time.monotonic()
    with _lock:
        for key, job in list(_jobs.items()):
            if now - job["created"] > _JOB_TTL_SECONDS:
                _jobs.pop(key, None)


def _extract_docx(data: bytes) -> str:
    # ZIP-bomb guard must precede python-docx parsing.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 400 or sum(z.file_size for z in entries) > _MAX_UNCOMPRESSED_ZIP_BYTES:
            raise ValueError("This Word file is too large to review safely.")
        if not any(z.filename == "word/document.xml" for z in entries):
            raise ValueError("This is not a valid Word document.")
    from docx import Document
    doc = Document(io.BytesIO(data))
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
    text = "\n".join(chunks)
    if not text.strip():
        raise ValueError("No readable text was found in the Word document.")
    return text


def _submitted_document():
    if request.content_length and request.content_length > _MAX_REQUEST_BYTES:
        raise ValueError("Submission is too large (2.5 MB limit).")
    if request.is_json:
        data = request.get_json(silent=True)
        if not isinstance(data, dict):
            raise ValueError("A JSON object is required.")
        document = data.get("document", "")
        goal = data.get("goal", "")
    else:
        document = request.form.get("document", "")
        goal = request.form.get("goal", "")
        uploaded = request.files.get("file")
        if uploaded and uploaded.filename:
            filename = uploaded.filename.lower()
            if not filename.endswith((".docx", ".txt")):
                raise ValueError("Upload a .docx or .txt document.")
            data = uploaded.read(_MAX_FILE_BYTES + 1)
            if len(data) > _MAX_FILE_BYTES:
                raise ValueError("The uploaded file exceeds the 2 MB limit.")
            try:
                decoded = _extract_docx(data) if filename.endswith(".docx") else data.decode("utf-8-sig")
            except (zipfile.BadZipFile, UnicodeDecodeError, OSError):
                raise ValueError("This document cannot be read. Use DOCX or UTF-8 TXT.")
            document = (str(document) + "\n\n" if document else "") + decoded
    if not isinstance(document, str) or not document.strip():
        raise ValueError("Paste text or choose a Word/TXT document.")
    if len(document) > MAX_DOCUMENT_CHARS:
        raise ValueError(f"Too much text: maximum {MAX_DOCUMENT_CHARS:,} characters.")
    if not isinstance(goal, str) or len(goal) > 2000:
        raise ValueError("Review instructions exceed 2,000 characters.")
    return document, (goal.strip() or "Review and improve this material; explain revisions and flag claims requiring verification.")


def _run_job(job_id: str, document: str, goal: str):
    with _lock:
        job = _jobs.get(job_id)
        if not job:
            return
        job["state"] = "running"
    try:
        # This invokes the two real independent AI lanes. No test fixtures.
        result = review_document(document, goal, options={"max_tokens": 3600})
        with _lock:
            job = _jobs.get(job_id)
            if job:
                job["result"] = result
                job["state"] = "finished"
    except Exception:
        # Never log customer input or provider error text.
        with _lock:
            job = _jobs.get(job_id)
            if job:
                job["state"] = "failed"
                job["error"] = "Review failed. Please try again later."


@bp.before_request
def protect_portal():
    g.macie_c_nonce = secrets.token_urlsafe(18)
    if not _is_ready():
        return jsonify({"error": "MACIE C login is not yet configured."}), 503


@bp.after_request
def security_headers(response):
    response.headers["Cache-Control"] = "no-store, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    response.headers["Content-Security-Policy"] = (
        "default-src 'none'; base-uri 'none'; frame-ancestors 'none'; "
        "form-action 'self'; connect-src 'self'; "
        "style-src 'unsafe-inline'; script-src 'nonce-" + g.macie_c_nonce + "'"
    )
    return response


@bp.route("", methods=["GET"])
@bp.route("/", methods=["GET"])
def home():
    email = _current_user()
    return make_response(render_template_string(
        PAGE_HTML,
        email=email,
        csrf=session.get("macie_c_csrf", "") if email else "",
        nonce=g.macie_c_nonce,
    ))


@bp.route("/login", methods=["POST"])
def login():
    if not _same_origin():
        return jsonify({"error": "Invalid request origin."}), 403
    if request.content_length and request.content_length > 2048:
        return jsonify({"error": "Invalid login."}), 400
    data = request.get_json(silent=True) or {}
    if not isinstance(data, dict):
        return jsonify({"error": "Invalid login."}), 400
    email = str(data.get("email", "")).strip().lower()[:200]
    password = data.get("password", "")
    if not isinstance(password, str) or len(password) > 250:
        return jsonify({"error": "Invalid credentials."}), 401
    key = f"{request.remote_addr}|{email}"
    if _login_blocked(key):
        return jsonify({"error": "Too many attempts. Please wait 15 minutes."}), 429
    record = _users().get(email, "")
    if not record or not _check_password(password, record):
        _record_login_failure(key)
        return jsonify({"error": "Invalid email or password."}), 401
    with _lock:
        _failed_logins.pop(key, None)
    session.clear()
    session.permanent = True
    session["macie_c_email"] = email
    session["macie_c_csrf"] = secrets.token_urlsafe(32)
    return jsonify({"ok": True, "email": email})


@bp.route("/logout", methods=["POST"])
def logout():
    denied = _authenticated()
    if denied:
        return denied
    if not _csrf_ok():
        return jsonify({"error": "Invalid CSRF token."}), 403
    session.clear()
    return jsonify({"ok": True})


@bp.route("/review", methods=["POST"])
def submit_review():
    denied = _authenticated()
    if denied:
        return denied
    if not _csrf_ok():
        return jsonify({"error": "Invalid CSRF token."}), 403
    try:
        document, goal = _submitted_document()
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    _cleanup_jobs()
    owner = _current_user()
    with _lock:
        user_running = sum(
            j["owner"] == owner and j["state"] in ("queued", "running")
            for j in _jobs.values()
        )
        total_running = sum(
            j["state"] in ("queued", "running") for j in _jobs.values()
        )
        if user_running >= 1 or total_running >= 3:
            return jsonify({"error": "A review is already processing. Wait for it to finish."}), 429
        job_id = secrets.token_urlsafe(18)
        _jobs[job_id] = {
            "owner": owner,
            "state": "queued",
            "created": time.monotonic(),
            "result": None,
        }
    _EXECUTOR.submit(_run_job, job_id, document, goal)
    return jsonify({"job_id": job_id, "status": "queued"}), 202


@bp.route("/review/<job_id>", methods=["GET"])
def review_status(job_id):
    denied = _authenticated()
    if denied:
        return denied
    _cleanup_jobs()
    with _lock:
        job = _jobs.get(job_id)
        if not job or job["owner"] != _current_user():
            return jsonify({"error": "Review not found or expired."}), 404
        response = {"status": job["state"]}
        if job["state"] == "finished":
            response["result"] = job["result"]
        if job["state"] == "failed":
            response["error"] = job.get("error", "Review failed.")
    return jsonify(response)
